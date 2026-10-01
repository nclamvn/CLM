#!/usr/bin/env python3
"""nap_tra_cong.py · Dua ket qua NGUOI TRA CONG vao truong ma_so_thue.

VI SAO CO (01/10/2026): lo dinh danh 02, anh Lam chot duong "ca hai". May gom ma so TU KHAI tu
website chinh chu (truong ma_so_tu_khai, khong tinh la da dinh danh). Nguoi cua RtR tra xac nhan
tren Cong thong tin quoc gia ve dang ky doanh nghiep (cong co captcha, may khong vuot), dien
PHIEU_TRA_CONG.xlsx va luu trang ket qua thanh PDF trong <dot>/tra_cong/. Script nay doc hai thu
do va CHI nap mot dong khi bang chung tu cong noi dung dieu nguoi tra ghi.

LUAT, moi dong co ma o cot H (Ma so tren cong):
  MA_SAI_DANG     ma khong dung dang 10 so hoac 10 so-3 so.
  THIEU_NGUOI     thieu nguoi tra (L) hoac ngay tra (M, dd/mm/yyyy).
  THIEU_PDF       file PDF o cot N khong co trong tra_cong/.
  PDF_KHONG_CONG  PDF khong mang dau cua cong (ten mien dangkykinhdoanh.gov.vn hoac ten cong).
  PDF_KHONG_MA    ma o cot H khong co trong chu cua PDF.
  PDF_KHONG_TEN   ten doanh nghiep o cot I khong co trong chu cua PDF.
  LECH_KHONG_GHI  ma cong khac ma tu khai ma cot O (ghi chu) de trong: lech phai co loi giai thich.
  DON_VI_LA       don vi o cot B khong co trong registry.
  MA_HAI_DON_VI   mot ma cong gan cho hai don vi (bai hoc HTI: me va con la hai phap nhan).
Mot vi pham bat ky -> khong nap gi ca (ca lo, khong nap nua voi).

Moi dong dat thanh mot claim ma_so_thue: tier A, extraction verbatim, evidence_span la doan chu
PDF quanh ma so, capture.source = dangkykinhdoanh.gov.vn (nam trong NGUON_CHINH_THUC cua
check_ma_so_thue.py, cong do giu nguyen). Ban chu cua PDF luu thanh ban chup .md co dong
"# TRANG TINH:" (trang tra cuu khong co ngay dang; hau to la ngay tra), PDF goc chep kem ben canh.

Phan logic (kiem_dong) la ham thuan, khong can openpyxl hay pdftotext: rang bite_nap_tra_cong.py
goi thang no, nen chay duoc ca trong CI. Doc xlsx va PDF chi o lop ngoai (main).

Chay: python3 nap_tra_cong.py <dot> [--ghi]    Mac dinh chay thu (in ra, khong ghi).
Exit 0 xong · 2 vi pham hoac tu choi · 3 KHONG CHAY DUOC.
"""
import json
import re
import subprocess
import sys
import unicodedata
from pathlib import Path

HERE = Path(__file__).resolve().parent
DOMAIN = HERE.parent / "domains" / "don_vi_cncl"
REG_WEB = HERE.parents[1] / ".touch" / "lib" / "cncl-registry.json"
CONG_URL = "https://dangkykinhdoanh.gov.vn/vn/Pages/Trangchu.aspx"
DAU_CONG = ("dangkykinhdoanh.gov.vn", "cổng thông tin quốc gia về đăng ký doanh nghiệp")
MAU_MA = re.compile(r"^\d{10}(-\d{3})?$")
MAU_NGAY = re.compile(r"^(\d{2})/(\d{2})/(\d{4})$")


def chuan(s):
    s = unicodedata.normalize("NFC", str(s or "")).replace("﻿", "")
    return re.sub(r"\s+", " ", s).strip()


def kiem_dong(dong, doc_pdf, ten_registry):
    """dong: [{stt, don_vi, ma_tu_khai, ma_cong, ten_cong, tinh_trang, nguoi, ngay, pdf, ghi_chu}]
    doc_pdf(ten_file) -> chu cua PDF hoac None neu khong co file.
    Tra (claims_du_kien, vi_pham)."""
    vi, ra, ma_cua = [], [], {}
    for d in dong:
        ma = chuan(d.get("ma_cong"))
        if not ma:
            continue
        vt = f"dong {d.get('stt')} ({d.get('don_vi')})"
        if d.get("don_vi") not in ten_registry:
            vi.append(f"DON_VI_LA: {vt}")
        if not MAU_MA.match(ma):
            vi.append(f"MA_SAI_DANG: {vt} '{ma}'")
        if not chuan(d.get("nguoi")) or not MAU_NGAY.match(chuan(d.get("ngay"))):
            vi.append(f"THIEU_NGUOI: {vt} can nguoi tra va ngay tra dd/mm/yyyy")
        tk = chuan(d.get("ma_tu_khai"))
        if tk and tk != ma and not chuan(d.get("ghi_chu")):
            vi.append(f"LECH_KHONG_GHI: {vt} ma cong {ma} khac ma tu khai {tk} ma khong co ghi chu")
        chu = doc_pdf(chuan(d.get("pdf")))
        if chu is None:
            vi.append(f"THIEU_PDF: {vt} khong co {d.get('pdf')}")
            continue
        c = chuan(chu)
        if not any(x in c.lower() for x in DAU_CONG):
            vi.append(f"PDF_KHONG_CONG: {vt} PDF khong mang dau cua cong dang ky doanh nghiep")
        if ma not in c:
            vi.append(f"PDF_KHONG_MA: {vt} ma {ma} khong co trong PDF")
        ten = chuan(d.get("ten_cong"))
        if not ten or ten.lower() not in c.lower():
            vi.append(f"PDF_KHONG_TEN: {vt} ten '{ten}' khong co trong PDF")
        ma_cua.setdefault(ma, set()).add(d.get("don_vi"))
        i = c.find(ma)
        span = c[max(0, i - 160): i + len(ma) + 160].strip() if i >= 0 else ""
        ra.append({"dong": d, "ma": ma, "span": span})
    for ma, ds in ma_cua.items():
        if len(ds) > 1:
            vi.append(f"MA_HAI_DON_VI: {ma} gan cho {' | '.join(sorted(map(str, ds)))}")
    return ra, vi


def doc_phieu(p):
    from openpyxl import load_workbook  # chi lop ngoai can; CI khong chay ham nay
    ws = load_workbook(p, data_only=True)["Phiếu tra"]
    dong = []
    for r in range(4, ws.max_row + 1):
        g = lambda c: ws.cell(row=r, column=c).value  # noqa: E731
        if g(1) is None:
            continue
        ngay = g(13)
        if hasattr(ngay, "strftime"):
            ngay = ngay.strftime("%d/%m/%Y")
        dong.append({"stt": g(1), "don_vi": g(2), "ma_tu_khai": g(4) or "", "ma_cong": g(8) or "",
                     "ten_cong": g(9) or "", "tinh_trang": g(10) or "", "nguoi": g(12) or "",
                     "ngay": ngay or "", "pdf": g(14) or "", "ghi_chu": g(15) or ""})
    return dong


def main(argv):
    a = argv[1:]
    if not a:
        print("KHONG CHAY DUOC: dung nap_tra_cong.py <dot> [--ghi]")
        return 3
    dot, ghi = Path(a[0]).resolve(), "--ghi" in a
    phieu, thu = dot / "PHIEU_TRA_CONG.xlsx", dot / "tra_cong"
    if not phieu.exists():
        print(f"KHONG CHAY DUOC: thieu {phieu}")
        return 3

    def doc_pdf(ten):
        p = thu / ten
        if not ten or not p.exists():
            return None
        r = subprocess.run(["pdftotext", "-layout", str(p), "-"], capture_output=True, text=True)
        return r.stdout if r.returncode == 0 else ""

    ten_reg = {u["name"] for u in json.loads(REG_WEB.read_text(encoding="utf-8"))["units"]}
    ra, vi = kiem_dong(doc_phieu(phieu), doc_pdf, ten_reg)
    da_co = {json.loads(l)["entity"] for l in (DOMAIN / "claims.jsonl").read_text(encoding="utf-8").splitlines()
             if l.strip() and json.loads(l)["field"] == "ma_so_thue"}
    moi = [x for x in ra if x["dong"]["don_vi"] not in da_co]
    print(f"TRA CONG: {len(ra)} dong co ma cong · {len(moi)} moi · {len(ra) - len(moi)} da nap truoc")
    for x in moi:
        d = x["dong"]
        print(f"  {str(d['don_vi'])[:44]:44} {x['ma']}  {d['tinh_trang']}  tra boi {d['nguoi']} {d['ngay']}")
    if vi:
        print(f"\nFAIL: {len(vi)} vi pham, khong nap dong nao")
        for v in vi:
            print("  " + v)
        return 2
    if not ghi or not moi:
        print("\nCHAY THU: khong ghi. Them --ghi de nap." if moi else "\nKhong co dong moi de nap.")
        return 0

    snap = DOMAIN / "snapshots"
    with (DOMAIN / "claims.jsonl").open("a", encoding="utf-8") as g:
        for x in moi:
            d = x["dong"]
            dd, mm, yy = MAU_NGAY.match(chuan(d["ngay"])).groups()
            goc = re.sub(r"\.pdf$", "", str(d["pdf"]))
            ten_md = f"dangkykinhdoanh_{goc}_{yy}{mm}{dd}.md"
            (snap / f"dangkykinhdoanh_{goc}_{yy}{mm}{dd}.pdf").write_bytes((thu / d["pdf"]).read_bytes())
            (snap / ten_md).write_text(
                f"# SNAPSHOT · dangkykinhdoanh.gov.vn · captured {yy}-{mm}-{dd} via nguoi tra cong (PDF, pdftotext)\n"
                f"# URL: {CONG_URL}\n"
                f"# Nguoi tra: {d['nguoi']}, ngay {d['ngay']}. PDF goc nam canh file nay. Tinh trang tren cong: {d['tinh_trang']}.\n"
                f"# TRANG TINH: trang tra cuu doanh nghiep khong co ngay dang. Hau to ten file la NGAY TRA.\n\n"
                + doc_pdf(d["pdf"]), encoding="utf-8")
            g.write(json.dumps({
                "entity": d["don_vi"], "field": "ma_so_thue", "value": x["ma"], "evidence_span": x["span"],
                "extraction": "verbatim", "tier": "A",
                "capture": {"url": CONG_URL, "fetched_at": f"{yy}-{mm}-{dd}T00:00:00Z", "snapshot": ten_md,
                            "source": "dangkykinhdoanh.gov.vn"},
                "note": f"TRA CONG boi {d['nguoi']} ngay {d['ngay']} (lo dinh danh {dot.name}). Ten tren cong: {chuan(d['ten_cong'])}. "
                        f"Tinh trang: {d['tinh_trang']}. Ma tu khai: {chuan(d['ma_tu_khai']) or 'khong co'}."
                        + (f" Ghi chu nguoi tra: {chuan(d['ghi_chu'])}" if chuan(d['ghi_chu']) else ""),
            }, ensure_ascii=False) + "\n")
    ns = DOMAIN / "ngan_sach_ma_so_thue.txt"
    dong_ns = ns.read_text(encoding="utf-8").splitlines()
    cu = next(int(x.strip()) for x in dong_ns if x.strip() and not x.startswith("#"))
    moi_so = cu + len(moi)
    ns.write_text("\n".join(
        [x for x in dong_ns if x.strip() != str(cu)]
        + [f"# {moi[0]['dong']['ngay']} nap_tra_cong.py: +{len(moi)} ma tu cong ({dot.name}), {cu} -> {moi_so}.", str(moi_so)]) + "\n",
        encoding="utf-8")
    print(f"\nDA NAP {len(moi)} ma so thue tu cong. Ngan sach do phu {cu} -> {moi_so}.")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
